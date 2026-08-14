"""Shared utilities for the bounded iron-ore RAG evaluation.

The module intentionally uses the Python standard library for HTTP and scoring.
Only local source verification needs ``openpyxl`` and the ``pdftotext`` binary.
"""

from __future__ import annotations

import hashlib
import json
import math
import re
import subprocess
import time
import unicodedata
import urllib.parse
import urllib.request
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple


CUTOFFS = (1, 3, 5, 10)
# Keep ``~`` because it is semantically meaningful in numeric ranges such as
# ``1.00%~40.00%``; stripping a Markdown marker must not collapse the range.
_STRIP_RE = re.compile(r"[\s　*#`_]+")
_SAFE_LABEL_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]{0,63}$")


def normalize(text: str) -> str:
    """Normalize source and chunk text before stable substring matching."""

    value = unicodedata.normalize("NFKC", text or "")
    value = value.translate(str.maketrans({"–": "-", "—": "-", "−": "-", "～": "~"}))
    return _STRIP_RE.sub("", value)


def load_jsonl(path: Path) -> List[dict]:
    rows: List[dict] = []
    with path.open(encoding="utf-8") as handle:
        for lineno, raw in enumerate(handle, 1):
            line = raw.strip()
            if not line:
                continue
            try:
                value = json.loads(line)
            except json.JSONDecodeError as exc:
                raise ValueError(f"{path}:{lineno}: JSON parse failed: {exc}") from exc
            if not isinstance(value, dict):
                raise ValueError(f"{path}:{lineno}: each JSONL row must be an object")
            rows.append(value)
    return rows


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_jsonl(path: Path, rows: Iterable[Mapping[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    text = "".join(json.dumps(dict(row), ensure_ascii=False) + "\n" for row in rows)
    path.write_text(text, encoding="utf-8")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def validate_label(label: str) -> str:
    if not _SAFE_LABEL_RE.fullmatch(label):
        raise ValueError("label may contain only letters, digits, '.', '_' and '-'")
    return label


def percentile(values: Sequence[float], quantile: float) -> Optional[float]:
    if not values:
        return None
    ordered = sorted(values)
    index = max(0, min(len(ordered) - 1, math.ceil(quantile * len(ordered)) - 1))
    return ordered[index]


def mean(values: Iterable[Optional[float]]) -> Optional[float]:
    usable = [value for value in values if value is not None]
    return sum(usable) / len(usable) if usable else None


def anchor_rank(anchor: str, contexts: Sequence[str]) -> int:
    target = normalize(anchor)
    if not target:
        return -1
    for index, context in enumerate(contexts):
        if target in normalize(context):
            return index
    return -1


def load_corpus_manifest(path: Path, repo_root: Path) -> Tuple[dict, Dict[str, dict]]:
    manifest = json.loads(path.read_text(encoding="utf-8"))
    documents: Dict[str, dict] = {}
    for item in manifest.get("documents", []):
        doc = dict(item)
        source = Path(doc["path"])
        if not source.is_absolute():
            source = repo_root / source
        doc["resolved_path"] = str(source.resolve())
        documents[doc["id"]] = doc
    return manifest, documents


def extract_source_text(document: Mapping[str, Any]) -> str:
    path = Path(str(document["resolved_path"]))
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("openpyxl is required to verify XLSX anchors") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        parts: List[str] = []
        try:
            for sheet in workbook.worksheets:
                parts.append(f"\n# sheet:{sheet.title}\n")
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value is not None and str(cell.value).strip():
                            parts.append(f"{cell.coordinate}\t{cell.value}\n")
        finally:
            workbook.close()
        return "".join(parts)
    if suffix == ".pdf":
        completed = subprocess.run(
            ["pdftotext", "-layout", str(path), "-"],
            check=False,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
        if completed.returncode != 0:
            error = completed.stderr.decode("utf-8", errors="replace").strip()
            raise RuntimeError(f"pdftotext failed for {path}: {error}")
        return completed.stdout.decode("utf-8", errors="replace")
    return path.read_text(encoding="utf-8")


class ApiError(RuntimeError):
    """Raised for transport, HTTP, or wrapped business API failures."""


class ApiClient:
    def __init__(self, base: str, token: Optional[str] = None, timeout: int = 180):
        self.base = base.rstrip("/")
        self.token = token
        self.timeout = timeout

    def login(self, username: str, password: str) -> str:
        data = self.request_json(
            "/auth/login", method="POST", body={"username": username, "password": password}, token=False
        )
        self.token = data["token"]
        return self.token

    def _headers(self, token: bool = True) -> Dict[str, str]:
        headers = {"Accept": "application/json"}
        if token and self.token:
            headers["Authorization"] = self.token
        return headers

    def request_json(
        self,
        path: str,
        *,
        method: str = "GET",
        body: Optional[Mapping[str, Any]] = None,
        query: Optional[Mapping[str, Any]] = None,
        token: bool = True,
    ) -> Any:
        url = self.base + path
        if query:
            url += "?" + urllib.parse.urlencode(query)
        raw = json.dumps(body, ensure_ascii=False).encode("utf-8") if body is not None else None
        request = urllib.request.Request(url, data=raw, method=method)
        for key, value in self._headers(token).items():
            request.add_header(key, value)
        if body is not None:
            request.add_header("Content-Type", "application/json; charset=utf-8")
        try:
            with urllib.request.urlopen(request, timeout=self.timeout) as response:
                payload = json.loads(response.read().decode("utf-8"))
        except Exception as exc:
            raise ApiError(f"{method} {url} failed: {exc}") from exc
        if not payload.get("success"):
            raise ApiError(f"{method} {url}: {payload.get('code')} {payload.get('message')}")
        return payload.get("data")

    def upload_file(self, path: str, fields: Mapping[str, str], file_path: Path) -> Any:
        boundary = f"----ragent-eval-{time.time_ns()}"
        chunks: List[bytes] = []
        for name, value in fields.items():
            chunks.extend(
                [
                    f"--{boundary}\r\n".encode(),
                    f'Content-Disposition: form-data; name="{name}"\r\n\r\n'.encode(),
                    str(value).encode("utf-8"),
                    b"\r\n",
                ]
            )
        chunks.extend(
            [
                f"--{boundary}\r\n".encode(),
                (
                    f'Content-Disposition: form-data; name="file"; '
                    f'filename="{file_path.name}"\r\n'
                ).encode("utf-8"),
                b"Content-Type: application/octet-stream\r\n\r\n",
                file_path.read_bytes(),
                b"\r\n",
                f"--{boundary}--\r\n".encode(),
            ]
        )
        request = urllib.request.Request(self.base + path, data=b"".join(chunks), method="POST")
        for key, value in self._headers().items():
            request.add_header(key, value)
        request.add_header("Content-Type", f"multipart/form-data; boundary={boundary}")
        try:
            with urllib.request.urlopen(request, timeout=self.timeout) as response:
                payload = json.loads(response.read().decode("utf-8"))
        except Exception as exc:
            raise ApiError(f"upload {file_path} failed: {exc}") from exc
        if not payload.get("success"):
            raise ApiError(f"upload {file_path}: {payload.get('code')} {payload.get('message')}")
        return payload.get("data")

    def query_eval(
        self,
        question: str,
        sub_questions: Optional[Sequence[str]] = None,
    ) -> Tuple[dict, int]:
        started = time.monotonic()
        if sub_questions is None:
            result = self.request_json("/rag/eval", query={"question": question})
        else:
            result = self.request_json(
                "/rag/eval/replay",
                method="POST",
                body={"question": question, "subQuestions": list(sub_questions)},
            )
        if not isinstance(result, dict):
            raise ApiError("eval endpoint returned a non-object response")
        return result, round((time.monotonic() - started) * 1000)

    def query_answer(self, question: str, deep_thinking: bool = False) -> dict:
        query = {"question": question, "deepThinking": str(deep_thinking).lower()}
        url = self.base + "/rag/v3/chat?" + urllib.parse.urlencode(query)
        request = urllib.request.Request(url, method="GET")
        request.add_header("Accept", "text/event-stream")
        if self.token:
            request.add_header("Authorization", self.token)
        started = time.monotonic()
        events: List[dict] = []
        answer: List[str] = []
        thinking: List[str] = []
        finish: Optional[dict] = None
        event_name = "message"
        data_lines: List[str] = []

        def dispatch() -> None:
            nonlocal event_name, data_lines, finish
            if not data_lines:
                event_name = "message"
                return
            raw_data = "\n".join(data_lines)
            try:
                payload: Any = json.loads(raw_data)
            except json.JSONDecodeError:
                payload = raw_data
            events.append({"event": event_name, "data": payload})
            if event_name == "message" and isinstance(payload, dict):
                if payload.get("type") == "response":
                    answer.append(str(payload.get("delta") or ""))
                elif payload.get("type") == "think":
                    thinking.append(str(payload.get("delta") or ""))
            elif event_name == "reject" and isinstance(payload, dict):
                answer.append(str(payload.get("delta") or ""))
            elif event_name in {"finish", "cancel"} and isinstance(payload, dict):
                finish = payload
            event_name = "message"
            data_lines = []

        try:
            with urllib.request.urlopen(request, timeout=self.timeout) as response:
                for raw_line in response:
                    line = raw_line.decode("utf-8", errors="replace").rstrip("\r\n")
                    if not line:
                        dispatch()
                    elif line.startswith(":"):
                        continue
                    elif line.startswith("event:"):
                        event_name = line[6:].strip()
                    elif line.startswith("data:"):
                        data_lines.append(line[5:].strip())
                dispatch()
        except Exception as exc:
            raise ApiError(f"SSE GET {url} failed after {len(events)} events: {exc}") from exc

        return {
            "answer": "".join(answer),
            "thinking": "".join(thinking),
            "finish": finish,
            "events": events,
            "wall_ms": round((time.monotonic() - started) * 1000),
        }


def validate_retrieval_diagnostics(
    response: Mapping[str, Any],
    refill_mode: str,
) -> Dict[str, Any]:
    """Validate the request-level fair-refill diagnostics returned by the eval API.

    The runner treats these fields as evidence rather than optional logging.  A
    mislabeled server (for example, a report declared as ``on`` while the
    backend still runs the control path) must fail before it can be compared.
    """

    if refill_mode not in {"off", "on"}:
        raise ValueError(f"unsupported refill mode: {refill_mode}")
    raw = response.get("retrievalDiagnostics")
    if not isinstance(raw, Mapping):
        raise ValueError("response is missing retrievalDiagnostics")

    expected_enabled = refill_mode == "on"
    enabled = raw.get("fairRefillEnabled")
    if not isinstance(enabled, bool):
        raise ValueError("retrievalDiagnostics.fairRefillEnabled must be boolean")
    if enabled != expected_enabled:
        raise ValueError(
            "retrievalDiagnostics.fairRefillEnabled does not match "
            f"--refill-mode {refill_mode}"
        )

    def integer(name: str) -> int:
        value = raw.get(name)
        if isinstance(value, bool) or not isinstance(value, int):
            raise ValueError(f"retrievalDiagnostics.{name} must be an integer")
        if value < 0:
            raise ValueError(f"retrievalDiagnostics.{name} must be >= 0")
        return value

    request_top_k = integer("requestTopK")
    candidate_count = integer("candidateCount")
    candidate_unique = integer("candidateUniqueCount")
    unique_before = integer("uniqueBeforeRefill")
    refill_added = integer("refillAdded")
    final_unique = integer("finalUniqueCount")
    unfilled_slots = integer("unfilledSlots")

    budgets = raw.get("initialBudgets")
    if not isinstance(budgets, list) or not budgets:
        raise ValueError("retrievalDiagnostics.initialBudgets must be a non-empty list")
    if any(isinstance(value, bool) or not isinstance(value, int) or value < 0 for value in budgets):
        raise ValueError("retrievalDiagnostics.initialBudgets must contain non-negative integers")

    if request_top_k <= 0:
        raise ValueError("retrievalDiagnostics.requestTopK must be > 0")
    if sum(budgets) != request_top_k:
        raise ValueError("retrievalDiagnostics.initialBudgets must sum to requestTopK")
    if candidate_unique > candidate_count:
        raise ValueError("candidateUniqueCount cannot exceed candidateCount")
    if unique_before > candidate_unique:
        raise ValueError("uniqueBeforeRefill cannot exceed candidateUniqueCount")
    if final_unique > candidate_unique:
        raise ValueError("finalUniqueCount cannot exceed candidateUniqueCount")
    if final_unique > request_top_k:
        raise ValueError("finalUniqueCount cannot exceed requestTopK")
    if final_unique != unique_before + refill_added:
        raise ValueError("finalUniqueCount must equal uniqueBeforeRefill + refillAdded")
    if unfilled_slots != request_top_k - final_unique:
        raise ValueError("unfilledSlots must equal requestTopK - finalUniqueCount")
    if not expected_enabled and refill_added != 0:
        raise ValueError("refillAdded must be zero while fair refill is disabled")

    chunk_ids = response.get("retrievedChunkIds")
    contexts = response.get("retrievedContexts")
    if not isinstance(chunk_ids, list) or not isinstance(contexts, list):
        raise ValueError("retrievedChunkIds and retrievedContexts must be lists")
    if any(not isinstance(chunk_id, str) or not chunk_id for chunk_id in chunk_ids):
        raise ValueError("retrievedChunkIds must contain non-empty strings")
    if len(chunk_ids) != len(set(chunk_ids)):
        raise ValueError("retrievedChunkIds must be unique")
    if len(chunk_ids) != final_unique or len(contexts) != final_unique:
        raise ValueError("finalUniqueCount must match returned chunk ids and contexts")

    return dict(raw)


def score_retrieval(
    row: Mapping[str, Any],
    response: Mapping[str, Any],
    doc_to_kb: Mapping[str, str],
    intent_mode: str,
) -> dict:
    contexts = list(response.get("retrievedContexts") or [])
    docs = list(response.get("retrievedDocIds") or [])
    context_docs = list(response.get("retrievedContextDocIds") or [])
    reference_docs = list(row.get("reference_docs") or [])
    anchors = list(row.get("reference_anchors") or [])
    score: Dict[str, Any] = {
        "n_chunks": len(contexts),
        "latency_ms": response.get("latencyMs"),
        "has_kb": bool(response.get("hasKb")),
        "retrieved_doc_count": len(set(docs)),
    }

    if row.get("answerable"):
        doc_set = set(docs)
        reference_set = set(reference_docs)
        hits = doc_set & reference_set
        score["doc_recall"] = len(hits) / len(reference_set) if reference_set else None
        score["doc_precision"] = len(hits) / len(doc_set) if doc_set else 0.0
        ranks = [anchor_rank(anchor, contexts) for anchor in anchors]
        hit_ranks = [rank for rank in ranks if rank >= 0]
        score["anchor_recall"] = len(hit_ranks) / len(anchors) if anchors else None
        score["anchor_mrr"] = 1.0 / (min(hit_ranks) + 1) if hit_ranks else 0.0
        score["missed_anchors"] = [anchor for anchor, rank in zip(anchors, ranks) if rank < 0]
        for cutoff in CUTOFFS:
            score[f"anchor_hit@{cutoff}_any"] = 1.0 if any(rank < cutoff for rank in hit_ranks) else 0.0
            score[f"anchor_hit@{cutoff}_all"] = (
                1.0 if anchors and all(0 <= rank < cutoff for rank in ranks) else 0.0
            )
        normalized_anchors = [normalize(anchor) for anchor in anchors if normalize(anchor)]
        useful = sum(
            1 for context in contexts if any(anchor in normalize(context) for anchor in normalized_anchors)
        )
        score["context_precision"] = useful / len(contexts) if contexts else 0.0
    else:
        score["negative_doc_spread"] = len(set(docs))

    if row.get("routing_scored"):
        expected = set(row.get("expected_kbs") or [])
        known_groups = [doc_to_kb[doc] for doc in context_docs if doc in doc_to_kb]
        correct = sum(1 for group in known_groups if group in expected)
        score["routing_purity"] = correct / len(known_groups) if known_groups else 0.0

    if intent_mode == "on" and row.get("intent_scored"):
        expected_ids = set(row.get("expected_intent_ids") or [])
        predicted = [item for item in response.get("intentLeafIds") or [] if item]
        score["intent_top1_correct"] = 1.0 if predicted and all(item in expected_ids for item in predicted) else 0.0
        score["predicted_intent_ids"] = predicted
    return score


def _metric_block(details: Sequence[Mapping[str, Any]]) -> Optional[dict]:
    if not details:
        return None
    metrics = [item["score"] for item in details]
    keys = [
        "doc_recall",
        "doc_precision",
        "anchor_recall",
        "anchor_mrr",
        "anchor_hit@5_any",
        "anchor_hit@5_all",
        "context_precision",
        "routing_purity",
        "intent_top1_correct",
    ]
    block = {"n": len(details)}
    for key in keys:
        value = mean(metric.get(key) for metric in metrics)
        if value is not None:
            block[key] = value
    return block


def aggregate_retrieval(details: Sequence[Mapping[str, Any]]) -> dict:
    positive = [item for item in details if item.get("answerable")]
    negative = [item for item in details if not item.get("answerable")]
    summary: Dict[str, Any] = {"overall_answerable": _metric_block(positive)}
    summary["by_family"] = {
        family: _metric_block([item for item in details if item.get("family") == family])
        for family in sorted({str(item.get("family")) for item in details})
    }
    summary["by_tier"] = {
        tier: _metric_block([item for item in details if item.get("tier") == tier])
        for tier in sorted({str(item.get("tier")) for item in details})
    }
    if negative:
        summary["unanswerable"] = {
            "n": len(negative),
            "avg_retrieved_doc_count": mean(item["score"].get("negative_doc_spread") for item in negative),
            "note": "retrieval alone cannot decide correct refusal; use the blinded answer review",
        }
    latency = [
        float(item["score"]["latency_ms"])
        for item in details
        if item["score"].get("latency_ms") is not None
    ]
    if latency:
        summary["latency_ms"] = {
            "mean": round(mean(latency) or 0),
            "p50": round(percentile(latency, 0.50) or 0),
            "p95": round(percentile(latency, 0.95) or 0),
            "max": round(max(latency)),
        }
    return summary
