"""Shared helpers for the knowledge-quality (kq) evaluation.

Standard library only for HTTP, scoring and text normalization. Source
verification needs ``pdftotext`` (poppler) for PDFs and ``openpyxl`` for XLSX.
Python 3.8 compatible.
"""

from __future__ import annotations

import hashlib
import json
import math
import re
import subprocess
import time
import unicodedata
import urllib.parse
import urllib.request
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple


CUTOFFS = (1, 3, 5, 10)
QUESTION_TYPES = ("numeric", "procedure", "survey", "confusion", "unanswerable")
SPLITS = ("tune", "test")

# ---------------------------------------------------------------------------
# Normalization (shared contract with the stage-2 Java normalizer via
# normalization_cases.json)
# ---------------------------------------------------------------------------

# Whitespace (ASCII, ideographic, zero-width, BOM) and Markdown emphasis marks are
# removed; ``~`` is kept because it carries range meaning (1.00%~10.00%).
_STRIP_RE = re.compile(r"[\s\u3000\u200b\u200c\u200d\ufeff*#`_]+")
# Applied before NFKC: NFKC would turn the ordinal indicator º into "o" and lose the degree meaning.
_PRE_MAP = str.maketrans({"º": "°", "˚": "°", "℃": "°C", "℉": "°F"})
# Applied after NFKC: dash variants that NFKC leaves alone.
_POST_MAP = str.maketrans({"–": "-", "—": "-", "−": "-", "―": "-", "‐": "-"})


def normalize(text: Optional[str]) -> str:
    """NFKC, unify dashes / degree signs / percent, drop whitespace and Markdown marks.

    Both the anchor and the chunk go through this before substring matching, so
    full-width digits, layout spaces inserted by ``pdftotext`` and the
    ``℃`` / ``°C`` split no longer break a match.
    """

    value = (text or "").translate(_PRE_MAP)
    value = unicodedata.normalize("NFKC", value)
    value = value.translate(_POST_MAP)
    return _STRIP_RE.sub("", value)


_MATH_SEGMENT_RE = re.compile(r"\$\$(.+?)\$\$|\$(.+?)\$|\\\((.+?)\\\)", re.S)
_MATH_WRAPPERS_RE = re.compile(r"\\(?:mathrm|mathit|mathbf|text|textrm|operatorname|mathsf|mathtt)\s*\{([^{}]*)\}")
_MATH_SIMPLE = (
    ("^{\\circ}", "°"), ("^\\circ", "°"), ("\\circ", "°"),
    ("\\pm", "±"), ("\\times", "×"), ("\\cdot", "·"), ("\\sim", "~"), ("\\approx", "≈"),
    ("\\%", "%"), ("\\leq", "≤"), ("\\le", "≤"), ("\\geq", "≥"), ("\\ge", "≥"),
    ("\\div", "÷"), ("\\mu", "μ"), ("\\rho", "ρ"), ("\\sigma", "σ"), ("\\beta", "β"), ("\\alpha", "α"),
)
_MATH_SPACING_RE = re.compile(r"\\[,;:!]|\\(?:quad|qquad|hspace\{[^}]*\})")
_MATH_LEFTOVER_CMD_RE = re.compile(r"\\[A-Za-z]+")


def _unwrap_math(segment: str) -> str:
    value = segment
    for _ in range(3):  # nested wrappers such as \mathrm{\mathrm{C}}
        value = _MATH_WRAPPERS_RE.sub(r"\1", value)
    value = _MATH_SPACING_RE.sub("", value)
    value = value.replace("~", "")  # non-breaking space inside math; must go before \sim becomes ~
    for src, dst in _MATH_SIMPLE:
        value = value.replace(src, dst)
    value = _MATH_LEFTOVER_CMD_RE.sub("", value)
    value = value.replace("{", "").replace("}", "").replace("^", "")
    return value


def strip_latex(text: Optional[str]) -> str:
    """Unwrap ``$...$`` math written by MinerU so ``$400 \\pm 20 ^{\\circ} \\mathrm{C}$`` reads ``400±20°C``.

    Text outside math segments is untouched. This feeds the "变形" column of the
    coverage metric; the raw column uses the chunk text as stored.
    """

    def repl(match: "re.Match[str]") -> str:
        body = next(group for group in match.groups() if group is not None)
        return _unwrap_math(body)

    return _MATH_SEGMENT_RE.sub(repl, text or "")


def normalize_latex(text: Optional[str]) -> str:
    return normalize(strip_latex(text))


# ---------------------------------------------------------------------------
# Character statistics
# ---------------------------------------------------------------------------

def count_digits(text: Optional[str]) -> int:
    """Count Unicode decimal digits (category Nd), which includes full-width ０-９."""

    return sum(1 for ch in text or "" if unicodedata.category(ch) == "Nd")


def count_ascii_digits(text: Optional[str]) -> int:
    return sum(1 for ch in text or "" if "0" <= ch <= "9")


def count_fullwidth_digits(text: Optional[str]) -> int:
    return sum(1 for ch in text or "" if "０" <= ch <= "９")


def count_pua(text: Optional[str]) -> int:
    """Private-use characters: glyphs without a Unicode mapping in the PDF font."""

    return sum(1 for ch in text or "" if unicodedata.category(ch) == "Co")


def count_non_space(text: Optional[str]) -> int:
    return sum(1 for ch in text or "" if not ch.isspace())


# Empty-slot patterns: a numeric token vanished and left its neighbours touching.
# Observed in the August ingestion of GB/T 6730.10: "温度控制在 ,放置 。", "硅含量高于 (质量分数)",
# "小于 时", "个国家的 个实验室".
SLOT_PATTERNS = [
    re.compile(pattern) for pattern in (
        r"在\s*[,，。.]",
        r"[(（]\s*见\s*[)）]",
        r"式\s*[(（]\s*[)）]",
        r"(?:小于|大于|等于|高于|低于|不大于|不小于|不超过|超过)\s*[时的,，(（]",
        r"的\s+个",
        r"(?:放置|保持|加热|灼烧|干燥|烘|称取|加入|加|移取|稀释至|冷却至)\s*[,，。.;；]",
        r"约\s*[,，。.]",
    )
]
NOISE_PATTERNS = [
    re.compile(pattern) for pattern in (
        r"中国标准出版社授权北京万方数据",
        r"推广使用",
        r"^\s*东北大学\s*$",
    )
]


def count_empty_slots(text: Optional[str]) -> int:
    value = text or ""
    return sum(len(pattern.findall(value)) for pattern in SLOT_PATTERNS)


def slots_per_1000(text: Optional[str]) -> Optional[float]:
    chars = count_non_space(text)
    if chars == 0:
        return None
    return count_empty_slots(text) * 1000.0 / chars


def count_noise_lines(text: Optional[str]) -> int:
    count = 0
    for line in (text or "").splitlines():
        if any(pattern.search(line) for pattern in NOISE_PATTERNS):
            count += 1
    return count


# ---------------------------------------------------------------------------
# Anchors
# ---------------------------------------------------------------------------

def contains_anchor(anchor: str, text: str, latex: bool = False) -> bool:
    target = normalize(anchor)
    if not target:
        return False
    haystack = normalize_latex(text) if latex else normalize(text)
    return target in haystack


def anchor_rank(anchor: str, contexts: Sequence[str], latex: bool = False) -> int:
    """Index of the first context containing the anchor, or -1."""

    target = normalize(anchor)
    if not target:
        return -1
    for index, context in enumerate(contexts):
        haystack = normalize_latex(context) if latex else normalize(context)
        if target in haystack:
            return index
    return -1


# ---------------------------------------------------------------------------
# Files
# ---------------------------------------------------------------------------

def utc_now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def load_jsonl(path: Path) -> List[dict]:
    rows: List[dict] = []
    with Path(path).open(encoding="utf-8") as handle:
        for lineno, raw in enumerate(handle, 1):
            line = raw.strip()
            if not line or line.startswith("//"):
                continue
            try:
                value = json.loads(line)
            except json.JSONDecodeError as exc:
                raise ValueError(f"{path}:{lineno}: JSON parse failed: {exc}") from exc
            if not isinstance(value, dict):
                raise ValueError(f"{path}:{lineno}: each JSONL row must be an object")
            rows.append(value)
    return rows


def write_jsonl(path: Path, rows: Iterable[Mapping[str, Any]]) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("".join(json.dumps(dict(row), ensure_ascii=False) + "\n" for row in rows), encoding="utf-8")


def read_json(path: Path) -> Any:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def write_json(path: Path, payload: Any) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def doc_stem(name: Optional[str]) -> str:
    """``铁矿石+硅含量的测定+重量法.pdf`` → ``铁矿石+硅含量的测定+重量法`` (matches the server's docName)."""

    if not name:
        return ""
    dot = name.rfind(".")
    return name[:dot] if 0 < dot < len(name) - 1 else name


def load_corpus(path: Path) -> dict:
    manifest = read_json(path)
    if not isinstance(manifest, dict) or not isinstance(manifest.get("documents"), list):
        raise ValueError(f"{path}: corpus manifest must have a documents list")
    return manifest


def corpus_documents(manifest: Mapping[str, Any]) -> Dict[str, dict]:
    return {doc["id"]: dict(doc) for doc in manifest.get("documents", [])}


def stem_to_doc_id(manifest: Mapping[str, Any]) -> Dict[str, str]:
    mapping: Dict[str, str] = {}
    for doc in manifest.get("documents", []):
        mapping[doc_stem(doc.get("doc_name") or Path(doc["path"]).name)] = doc["id"]
    return mapping


# ---------------------------------------------------------------------------
# Source text
# ---------------------------------------------------------------------------

def pdftotext_layout(path: Path) -> str:
    completed = subprocess.run(
        ["pdftotext", "-layout", str(path), "-"],
        check=False,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    if completed.returncode != 0:
        error = completed.stderr.decode("utf-8", errors="replace").strip()
        raise RuntimeError(f"pdftotext failed for {path}: {error}")
    return completed.stdout.decode("utf-8", errors="replace")


def xlsx_text(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise RuntimeError("openpyxl is required to read XLSX sources") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    parts: List[str] = []
    try:
        for sheet in workbook.worksheets:
            parts.append(f"\n# sheet:{sheet.title}\n")
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None and str(cell.value).strip():
                        parts.append(f"{cell.coordinate}\t{cell.value}\n")
    finally:
        workbook.close()
    return "".join(parts)


def pymupdf_text(path: Path) -> str:
    """Plain-text reading of the same text layer via PyMuPDF; empty when the module is missing."""

    try:
        import pymupdf  # type: ignore
    except ImportError:  # pragma: no cover - environment dependent
        return ""
    document = pymupdf.open(str(path))
    try:
        return "\n".join(page.get_text("text") for page in document)
    finally:
        document.close()


def extract_source_text(path: Path) -> str:
    """Source text for anchor verification.

    PDFs return the ``pdftotext -layout`` reading followed by the PyMuPDF reading of the same text
    layer: the layout mode keeps the E-BZ font digits on one line but can scramble the order of a
    split decimal in two-column regions (``1.\n00%~\n00%。\n15.``), while PyMuPDF keeps the order.
    An anchor counts as located when either reading contains it. Digit counts stay layout-based
    (pdf_textlayer.py), as the plan fixes the text-layer baseline on ``pdftotext -layout``.
    """

    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        layout = pdftotext_layout(path)
        plain = pymupdf_text(path)
        return layout + "\n\n" + plain if plain else layout
    if suffix in {".xlsx", ".xlsm"}:
        return xlsx_text(path)
    return path.read_text(encoding="utf-8")


# ---------------------------------------------------------------------------
# API client
# ---------------------------------------------------------------------------

class ApiError(RuntimeError):
    """Transport, HTTP or wrapped business failure of the Ragent API."""


class ApiClient:
    def __init__(self, base: str, token: Optional[str] = None, timeout: int = 180):
        self.base = base.rstrip("/")
        self.token = token
        self.timeout = timeout

    def login(self, username: str, password: str) -> str:
        data = self.request_json(
            "/auth/login", method="POST", body={"username": username, "password": password}, token=False
        )
        self.token = data["token"]
        return self.token

    def _headers(self, token: bool = True) -> Dict[str, str]:
        headers = {"Accept": "application/json"}
        if token and self.token:
            headers["Authorization"] = self.token
        return headers

    def request_json(
        self,
        path: str,
        *,
        method: str = "GET",
        body: Optional[Mapping[str, Any]] = None,
        query: Optional[Mapping[str, Any]] = None,
        token: bool = True,
    ) -> Any:
        url = self.base + path
        if query:
            url += "?" + urllib.parse.urlencode(query)
        raw = json.dumps(body, ensure_ascii=False).encode("utf-8") if body is not None else None
        request = urllib.request.Request(url, data=raw, method=method)
        for key, value in self._headers(token).items():
            request.add_header(key, value)
        if body is not None:
            request.add_header("Content-Type", "application/json; charset=utf-8")
        try:
            with urllib.request.urlopen(request, timeout=self.timeout) as response:
                payload = json.loads(response.read().decode("utf-8"))
        except Exception as exc:
            raise ApiError(f"{method} {url} failed: {exc}") from exc
        if not isinstance(payload, dict) or not payload.get("success"):
            code = payload.get("code") if isinstance(payload, dict) else None
            message = payload.get("message") if isinstance(payload, dict) else payload
            raise ApiError(f"{method} {url}: {code} {message}")
        return payload.get("data")

    def upload_file(self, path: str, fields: Mapping[str, str], file_path: Path) -> Any:
        boundary = f"----ragent-kq-{time.time_ns()}"
        chunks: List[bytes] = []
        for name, value in fields.items():
            chunks.extend([
                f"--{boundary}\r\n".encode(),
                f'Content-Disposition: form-data; name="{name}"\r\n\r\n'.encode(),
                str(value).encode("utf-8"),
                b"\r\n",
            ])
        chunks.extend([
            f"--{boundary}\r\n".encode(),
            f'Content-Disposition: form-data; name="file"; filename="{file_path.name}"\r\n'.encode("utf-8"),
            b"Content-Type: application/octet-stream\r\n\r\n",
            Path(file_path).read_bytes(),
            b"\r\n",
            f"--{boundary}--\r\n".encode(),
        ])
        request = urllib.request.Request(self.base + path, data=b"".join(chunks), method="POST")
        for key, value in self._headers().items():
            request.add_header(key, value)
        request.add_header("Content-Type", f"multipart/form-data; boundary={boundary}")
        try:
            with urllib.request.urlopen(request, timeout=self.timeout) as response:
                payload = json.loads(response.read().decode("utf-8"))
        except Exception as exc:
            raise ApiError(f"upload {file_path} failed: {exc}") from exc
        if not payload.get("success"):
            raise ApiError(f"upload {file_path}: {payload.get('code')} {payload.get('message')}")
        return payload.get("data")

    def replay(self, question: str, sub_questions: Optional[Sequence[str]] = None) -> Tuple[dict, int]:
        """``POST /rag/eval/replay``: live rewrite when ``sub_questions`` is None, replay otherwise."""

        body: Dict[str, Any] = {"question": question}
        if sub_questions is not None:
            body["subQuestions"] = list(sub_questions)
        started = time.monotonic()
        data = self.request_json("/rag/eval/replay", method="POST", body=body)
        if not isinstance(data, dict):
            raise ApiError("eval endpoint returned a non-object response")
        return data, round((time.monotonic() - started) * 1000)


# ---------------------------------------------------------------------------
# Scoring
# ---------------------------------------------------------------------------

def final_chunks(response: Mapping[str, Any]) -> List[dict]:
    """Request-level final context in prompt order, de-duplicated by chunk id."""

    picked: Dict[str, dict] = {}
    for result in response.get("results") or []:
        for candidate in result.get("candidates") or []:
            if not candidate.get("finalSelected"):
                continue
            chunk_id = candidate.get("id")
            if not chunk_id or chunk_id in picked:
                continue
            picked[chunk_id] = {
                "id": chunk_id,
                "docId": candidate.get("docId"),
                "docName": candidate.get("docName"),
                "text": candidate.get("text") or "",
                "channelScore": candidate.get("channelScore"),
                "rerankScore": candidate.get("rerankScore"),
                "finalRank": candidate.get("finalRank"),
                "subQuestion": result.get("subQuestion"),
            }
    ordered = sorted(picked.values(), key=lambda item: (item["finalRank"] is None, item["finalRank"] or 0))
    ids_in_order = [chunk_id for chunk_id in response.get("finalChunkIds") or [] if chunk_id in picked]
    if len(ids_in_order) == len(picked):
        return [picked[chunk_id] for chunk_id in ids_in_order]
    return ordered


def all_candidates(response: Mapping[str, Any]) -> List[dict]:
    rows: List[dict] = []
    for result in response.get("results") or []:
        for candidate in result.get("candidates") or []:
            rows.append(candidate)
    return rows


def max_rerank_score(response: Mapping[str, Any]) -> Optional[float]:
    scores = [c.get("rerankScore") for c in all_candidates(response) if c.get("rerankScore") is not None]
    return max(scores) if scores else None


def score_question(row: Mapping[str, Any], response: Mapping[str, Any], stem_to_id: Mapping[str, str]) -> dict:
    chunks = final_chunks(response)
    contexts = [chunk["text"] for chunk in chunks]
    doc_ids = [stem_to_id.get(doc_stem(chunk.get("docName")), doc_stem(chunk.get("docName"))) for chunk in chunks]
    top = chunks[0] if chunks else None
    score: Dict[str, Any] = {
        "n_chunks": len(chunks),
        "retrieved_doc_count": len(set(doc_ids)),
        "latency_ms": response.get("latencyMs"),
        "max_rerank_score": max_rerank_score(response),
        "top_final_rerank": None if top is None else top.get("rerankScore"),
        "sub_question_count": len(response.get("subQuestions") or []),
    }
    if not row.get("answerable"):
        return score

    reference = list(row.get("reference_docs") or [])
    anchors = list(row.get("anchors") or [])
    hits = set(doc_ids) & set(reference)
    score["doc_recall"] = len(hits) / len(reference) if reference else None
    ranks = [anchor_rank(anchor, contexts) for anchor in anchors]
    hit_ranks = [rank for rank in ranks if rank >= 0]
    score["anchor_recall"] = len(hit_ranks) / len(anchors) if anchors else None
    score["mrr"] = 1.0 / (min(hit_ranks) + 1) if hit_ranks else 0.0
    score["missed_anchors"] = [anchor for anchor, rank in zip(anchors, ranks) if rank < 0]
    for cutoff in CUTOFFS:
        score[f"hit@{cutoff}"] = 1.0 if any(rank < cutoff for rank in hit_ranks) else 0.0
        score[f"hit@{cutoff}_all"] = 1.0 if anchors and all(0 <= rank < cutoff for rank in ranks) else 0.0
    normalized_anchors = [normalize(anchor) for anchor in anchors if normalize(anchor)]
    useful = sum(1 for context in contexts if any(anchor in normalize(context) for anchor in normalized_anchors))
    score["context_precision"] = useful / len(contexts) if contexts else 0.0
    return score


ANSWERABLE_METRICS = ("hit@5", "hit@5_all", "anchor_recall", "mrr", "context_precision", "doc_recall", "n_chunks")
UNANSWERABLE_METRICS = ("n_chunks", "retrieved_doc_count", "max_rerank_score", "top_final_rerank")


def mean(values: Iterable[Optional[float]]) -> Optional[float]:
    usable = [float(value) for value in values if value is not None]
    return sum(usable) / len(usable) if usable else None


def percentile(values: Sequence[float], quantile: float) -> Optional[float]:
    if not values:
        return None
    ordered = sorted(values)
    index = max(0, min(len(ordered) - 1, math.ceil(quantile * len(ordered)) - 1))
    return ordered[index]


def _block(details: Sequence[Mapping[str, Any]], keys: Sequence[str]) -> Optional[dict]:
    if not details:
        return None
    block: Dict[str, Any] = {"n": len(details)}
    for key in keys:
        value = mean(item["score"].get(key) for item in details)
        if value is not None:
            block[key] = value
    return block


def aggregate(details: Sequence[Mapping[str, Any]]) -> dict:
    answerable = [item for item in details if item.get("answerable")]
    unanswerable = [item for item in details if not item.get("answerable")]
    summary: Dict[str, Any] = {
        "overall_answerable": _block(answerable, ANSWERABLE_METRICS),
        "by_type": {},
        "by_doc": {},
    }
    for qtype in sorted({str(item.get("type")) for item in answerable}):
        summary["by_type"][qtype] = _block([item for item in answerable if item.get("type") == qtype], ANSWERABLE_METRICS)
    for doc in sorted({ref for item in answerable for ref in item.get("reference_docs") or []}):
        summary["by_doc"][doc] = _block([item for item in answerable if doc in (item.get("reference_docs") or [])], ANSWERABLE_METRICS)
    if unanswerable:
        summary["unanswerable"] = _block(unanswerable, UNANSWERABLE_METRICS)
    latency = [float(item["score"]["latency_ms"]) for item in details if item["score"].get("latency_ms") is not None]
    if latency:
        summary["latency_ms"] = {
            "mean": round(mean(latency) or 0),
            "p50": round(percentile(latency, 0.50) or 0),
            "p95": round(percentile(latency, 0.95) or 0),
            "max": round(max(latency)),
        }
    return summary


def summarize_repeats(blocks: Sequence[Optional[Mapping[str, Any]]], keys: Sequence[str]) -> dict:
    """Mean and range (max − min) of each metric across repeated runs of one arm."""

    result: Dict[str, Any] = {"repeats": len(blocks)}
    for key in keys:
        values = [float(block[key]) for block in blocks if block and block.get(key) is not None]
        if not values:
            continue
        result[key] = {
            "values": values,
            "mean": sum(values) / len(values),
            "min": min(values),
            "max": max(values),
            "range": max(values) - min(values),
        }
    return result
